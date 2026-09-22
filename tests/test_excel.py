import numpy as np
from openpyxl import load_workbook

from motronic_autotuner.core.excel import export_maps
from motronic_autotuner.core.maps import CalibrationMap


def test_export_layout(tmp_path):
    values = np.array([[1.0, np.nan], [3.123456789, 4.0]])
    counts = np.array([[2.0, 0.0], [5.0, 1.0]])
    m1 = CalibrationMap("a", "MAP A", values, np.array([10, 20]), np.array([1000, 2000]),
                        counts=counts, counts_title="A - SAMPLE WEIGHTS",
                        secondary_x_label="Phys", secondary_x=np.array([250.04, 500.06]))
    m2 = CalibrationMap("b", "CURVE B", np.array([7.0, np.nan, 9.0]), np.array([1, 2, 3]), ["Trim %"])
    path = export_maps(tmp_path / "out.xlsx", [m1, m2])
    ws = load_workbook(path)["Tuning Maps"]

    assert ws["A1"].value == "MAP A"
    assert ws["F1"].value == "A - SAMPLE WEIGHTS"  # 2 x cols + label + 2 gap -> offset 5
    assert ws["A2"].value == "Y-Axis \\ X-Axis" and ws["B2"].value == 10 and ws["C2"].value == 20
    assert ws["A3"].value == "Phys" and ws["B3"].value == 250.0 and ws["C3"].value == 500.1
    assert ws["A4"].value == 1000 and ws["B4"].value == 1.0 and ws["C4"].value is None
    assert ws["B5"].value == 3.12346
    assert ws["G4"].value == 2.0 and ws["H4"].value == 0.0

    # blank row, then the 1D curve
    assert ws["A7"].value == "CURVE B"
    assert ws["A8"].value == "Y-Axis \\ X-Axis"
    assert ws["A9"].value == "Trim %" and ws["B9"].value == 7.0 and ws["C9"].value is None and ws["D9"].value == 9.0


def test_transposed_values_are_healed(tmp_path):
    m = CalibrationMap("t", "T", np.ones((3, 2)), np.array([1, 2, 3]), np.array([10, 20]))
    path = export_maps(tmp_path / "t.xlsx", [m])
    ws = load_workbook(path)["Tuning Maps"]
    assert ws["D3"].value == 1.0 and ws["A4"].value == 20 and ws["A5"].value is None
