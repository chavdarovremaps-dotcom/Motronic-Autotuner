"""Compare two 'Tuning Maps' workbooks block by block (values and sample weights).

Usage:
    python tools/compare_excel.py <a.xlsx> <b.xlsx> [--tol 1e-4]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

GAP_COLS = 2


def read_blocks(path: Path) -> dict[str, tuple[np.ndarray, np.ndarray | None]]:
    blocks = {}
    for ws in load_workbook(path, data_only=True).worksheets:
        blocks.update(_read_sheet([list(r) for r in ws.iter_rows(values_only=True)]))
    return blocks


def _read_sheet(rows):
    blocks = {}
    r = 0
    while r < len(rows):
        title = rows[r][0]
        if isinstance(title, str) and title and r + 1 < len(rows) and _is_corner(rows[r + 1][0]):
            header = rows[r + 1]
            nx = 0
            for v in header[1:]:
                if v is None or v == "":
                    break
                nx += 1
            counts_col = nx + 1 + GAP_COLS  # 0-based column of the counts label
            has_counts = len(rows[r]) > counts_col and rows[r][counts_col] not in (None, "")
            r += 2
            if isinstance(rows[r][0], str) and r + 1 < len(rows) and rows[r + 1][0] not in (None, ""):
                r += 1  # secondary header
            vals, cnts = [], []
            while r < len(rows) and rows[r][0] not in (None, ""):
                row = rows[r] + [None] * (counts_col + nx + 1 - len(rows[r]))
                vals.append([np.nan if v in (None, "") else float(v) for v in row[1:1 + nx]])
                if has_counts:
                    cnts.append([np.nan if v in (None, "") else float(v) for v in row[counts_col + 1:counts_col + 1 + nx]])
                r += 1
            blocks[title] = (np.array(vals, dtype=float), np.array(cnts, dtype=float) if has_counts else None)
        else:
            r += 1
    return blocks


def _is_corner(v) -> bool:
    """The header row's first cell: 'Y-Axis \\ X-Axis', 'MAP \\ RPM' and the like."""
    return isinstance(v, str) and "\\" in v


def compare(name: str, a: np.ndarray | None, b: np.ndarray | None, tol: float) -> bool:
    if a is None and b is None:
        return True
    if a is None or b is None:
        print(f"  {name:20s} present in only one file")
        return False
    if a.shape != b.shape:
        print(f"  {name:20s} shape {a.shape} vs {b.shape}")
        return False
    nan_mismatch = int((np.isnan(a) != np.isnan(b)).sum())
    both = ~np.isnan(a) & ~np.isnan(b)
    maxd = float(np.abs(a[both] - b[both]).max()) if both.any() else 0.0
    ok = nan_mismatch == 0 and maxd <= tol
    print(f"  {name:20s} shape {a.shape}  cells {int(both.sum()):4d}  max|diff| {maxd:.6f}  nan-mismatch {nan_mismatch}  {'OK' if ok else 'DIFF'}")
    if not ok:
        bad = np.argwhere((np.isnan(a) != np.isnan(b)) | (both & (np.abs(np.nan_to_num(a - b)) > tol)))
        for rr, cc in bad[:8]:
            print(f"      [{rr:2d},{cc:2d}] a={a[rr, cc]:12.6f} b={b[rr, cc]:12.6f}")
    return ok


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("a")
    ap.add_argument("b")
    ap.add_argument("--tol", type=float, default=1e-4)
    args = ap.parse_args()
    A, B = read_blocks(Path(args.a)), read_blocks(Path(args.b))
    all_ok = True
    for title in list(A) + [t for t in B if t not in A]:
        print(title)
        if title not in A or title not in B:
            print("  present in only one file")
            all_ok = False
            continue
        all_ok &= compare("values", A[title][0], B[title][0], args.tol)
        all_ok &= compare("sample weights", A[title][1], B[title][1], args.tol)
    print("\nRESULT:", "IDENTICAL within tolerance" if all_ok else "DIFFERENCES FOUND")
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())
