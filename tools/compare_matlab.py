"""Compare the Python pipeline against an Excel file the MATLAB app produced.

Usage:
    python tools/compare_matlab.py <log folder> <matlab xlsx> [--profile Used_ECU_Profile.json] [--param key=value ...]

Runs ingestion and every generator on the folder with the given profile,
then matches each calculated map to the block with the same title in the
MATLAB workbook and prints the maximum absolute difference.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from motronic_autotuner.core.preset import Preset  # noqa: E402
from motronic_autotuner.families.bosch_me7 import BOSCH_ME7  # noqa: E402


def read_blocks(path: Path) -> dict[str, np.ndarray]:
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    blocks: dict[str, np.ndarray] = {}
    r = 0
    while r < len(rows):
        title = rows[r][0]
        if isinstance(title, str) and title and isinstance(rows[r + 1][0], str) and "\\" in rows[r + 1][0]:
            header = rows[r + 1]
            # the counts twin shares this row further right; count only the contiguous x cells
            nx = 0
            for v in header[1:]:
                if v is None or v == "":
                    break
                nx += 1
            r += 2
            # optional secondary header: a text label in column A followed by another data row
            if isinstance(rows[r][0], str) and r + 1 < len(rows) and rows[r + 1][0] not in (None, ""):
                r += 1
            data = []
            while r < len(rows) and rows[r][0] is not None and rows[r][0] != "":
                vals = rows[r][1:1 + nx]
                data.append([np.nan if v in (None, "") else float(v) for v in vals])
                r += 1
            blocks[title] = np.array(data, dtype=float)
        else:
            r += 1
    return blocks


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("xlsx")
    ap.add_argument("--profile", default="Used_ECU_Profile.json")
    ap.add_argument("--param", action="append", default=[], help="override preset.params, e.g. min_samples=2")
    ap.add_argument("--detail", action="store_true", help="print differing cells")
    ap.add_argument("--dump", action="append", default=[], help="print ours and MATLAB for maps whose title contains this")
    args = ap.parse_args()

    folder = Path(args.folder)
    preset = Preset.load(folder / args.profile if not Path(args.profile).is_absolute() else args.profile)
    BOSCH_ME7.fill_defaults(preset)
    for kv in args.param:
        k, v = kv.split("=", 1)
        preset.params[k] = float(v) if v.replace(".", "", 1).lstrip("-").isdigit() else v
    print("params:", {k: v for k, v in preset.params.items() if k != "axis_wgdc_splat"})
    print("prep:", preset.prep)

    logs = BOSCH_ME7.ingest(folder, preset)
    print("logs:", logs.counts())
    result = BOSCH_ME7.run_all(preset, logs)
    for row in result.rows:
        print(f"  {row.name:55s} {row.status:30s} {row.size}")

    matlab = read_blocks(Path(args.xlsx))
    print(f"\nMATLAB blocks: {list(matlab)}\n")
    worst = 0.0
    for m in result.maps:
        ref = matlab.get(m.title)
        ours = m.as_2d()
        if ref is None:
            print(f"{m.title:55s} not in MATLAB file")
            continue
        if ref.shape != ours.shape:
            print(f"{m.title:55s} shape ours {ours.shape} vs matlab {ref.shape}")
            continue
        both = ~np.isnan(ref) & ~np.isnan(ours)
        nan_mismatch = int((np.isnan(ref) != np.isnan(ours)).sum())
        diff = np.abs(ref[both] - ours[both])
        maxd = float(diff.max()) if diff.size else 0.0
        worst = max(worst, maxd)
        print(f"{m.title:55s} max|diff| {maxd:10.5f}  cells {both.sum():4d}  nan-mismatch {nan_mismatch}")
        if any(s.lower() in m.title.lower() for s in args.dump):
            np.set_printoptions(linewidth=200, precision=4, suppress=True)
            print("  ours:\n", ours)
            print("  matlab:\n", ref)
            if m.counts is not None:
                print("  counts:\n", m.counts_2d())
        if args.detail and (maxd > 1e-4 or nan_mismatch):
            counts = m.counts_2d()
            bad = np.argwhere((np.isnan(ref) != np.isnan(ours)) | (both & (np.abs(np.nan_to_num(ref - ours)) > 1e-4)))
            for rr, cc in bad[:12]:
                cnt = f" count={counts[rr, cc]:.6f}" if counts is not None else ""
                print(f"      [{rr:2d},{cc:2d}] ours={ours[rr, cc]:12.6f} matlab={ref[rr, cc]:12.6f}{cnt}")
    print(f"\nworst max|diff| = {worst:.6f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
