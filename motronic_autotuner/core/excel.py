"""Excel export: every map as a titled block on one sheet, with its
sample-weight twin two columns to the right.

Port of Scripts/Utilities/ExportAllMapsToExcel.m, including its dimension
healing (transpose or pad a map whose shape disagrees with its axes).
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook

from .maps import CalibrationMap

GAP_COLS = 2
ROUND_DIGITS = 5


def _align(values: np.ndarray | None, ny: int, nx: int) -> np.ndarray | None:
    if values is None:
        return None
    v = np.atleast_2d(np.asarray(values, dtype=float))
    if v.shape == (ny, nx):
        return v
    if v.shape == (nx, ny):
        return v.T
    out = np.full((ny, nx), np.nan)
    r, c = min(v.shape[0], ny), min(v.shape[1], nx)
    out[:r, :c] = v[:r, :c]
    return out


def _cell(v) -> float | None:
    v = float(v)
    return None if np.isnan(v) else round(v, ROUND_DIGITS)


def export_maps(path: str | Path, maps: list[CalibrationMap], sheet: str = "Tuning Maps") -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    row = 1

    for m in maps:
        if m.is_empty:
            continue
        x_axis = np.asarray(m.x_axis, dtype=float).ravel()
        y_labels = m.y_labels()
        nx, ny = x_axis.size, len(y_labels)
        values = _align(m.values, ny, nx)
        counts = _align(m.counts, ny, nx)
        col_offset = nx + 1 + GAP_COLS

        ws.cell(row=row, column=1, value=m.title)
        if counts is not None:
            ws.cell(row=row, column=1 + col_offset, value=m.counts_title or f"{m.title} - SAMPLE WEIGHTS")
        row += 1

        _write_row(ws, row, 1, ["Y-Axis \\ X-Axis", *[float(v) for v in x_axis]])
        if counts is not None:
            _write_row(ws, row, 1 + col_offset, ["Y-Axis \\ X-Axis", *[float(v) for v in x_axis]])
        row += 1

        if m.secondary_x is not None and np.asarray(m.secondary_x).size:
            sec = [round(float(v), 1) for v in np.asarray(m.secondary_x).ravel()]
            _write_row(ws, row, 1, [m.secondary_x_label, *sec])
            if counts is not None:
                _write_row(ws, row, 1 + col_offset, [m.secondary_x_label, *sec])
            row += 1

        y_numeric = not isinstance(m.y_axis, list)
        for r in range(ny):
            label = float(np.asarray(m.y_axis).ravel()[r]) if y_numeric else y_labels[r]
            _write_row(ws, row, 1, [label, *[_cell(v) for v in values[r]]])
            if counts is not None:
                _write_row(ws, row, 1 + col_offset, [label, *[_cell(v) for v in counts[r]]])
            row += 1
        row += 1  # blank separator row

    wb.save(path)
    return path


def _write_row(ws, row: int, start_col: int, cells: list) -> None:
    for i, v in enumerate(cells):
        if v is not None:
            ws.cell(row=row, column=start_col + i, value=v)
